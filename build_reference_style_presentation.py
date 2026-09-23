from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUTPUT = "Praman_Setu_SIH26231_Reference_Style_Legal.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

NAVY = RGBColor(13, 34, 59)
BLUE = RGBColor(25, 104, 166)
CYAN = RGBColor(45, 183, 203)
GREEN = RGBColor(48, 151, 111)
ORANGE = RGBColor(238, 139, 67)
RED = RGBColor(190, 63, 68)
GOLD = RGBColor(242, 190, 70)
WHITE = RGBColor(255, 255, 255)
LIGHT = RGBColor(239, 246, 250)
INK = RGBColor(25, 42, 57)
MUTED = RGBColor(93, 113, 128)


def shape(slide, kind, x, y, width, height, fill, line=None):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(width), Inches(height))
    item.fill.solid()
    item.fill.fore_color.rgb = fill
    item.line.color.rgb = line or fill
    return item


def text(slide, value, x, y, width, height, size=14, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = valign
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return box


def hyperlink_text(slide, label, url, x, y, width, height, size=10):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = label
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = BLUE
    run.font.underline = True
    run.hyperlink.address = url
    return box


def header(slide, title, subtitle, number):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.18, BLUE)
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 7.2, 13.333, 0.3, NAVY)
    text(slide, "@ SIH IDEA SUBMISSION", 0.48, 0.28, 2.3, 0.2, 9, BLUE, True)
    text(slide, title, 0.48, 0.52, 8.7, 0.5, 27, NAVY, True, font="Aptos Display")
    text(slide, subtitle, 0.5, 1.07, 8.8, 0.28, 11, MUTED)
    text(slide, f"{number:02d}", 12.15, 0.39, 0.55, 0.3, 12, BLUE, True, align=PP_ALIGN.RIGHT)
    text(slide, "PRAMAN SETU  /  SIH26231", 0.5, 7.27, 2.4, 0.12, 7, WHITE, True)
    text(slide, "SIMULATED - NOT A DRUG TEST", 9.8, 7.27, 3.0, 0.12, 7, GOLD, True, align=PP_ALIGN.RIGHT)


def tag(slide, value, x, y, width, fill=CYAN):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, 0.31, fill, fill)
    text(slide, value, x, y + 0.03, width, 0.18, 8, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def card(slide, title, body, x, y, width, height, accent=BLUE):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, height, LIGHT, LIGHT)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.09, height, accent)
    text(slide, title, x + 0.22, y + 0.18, width - 0.4, 0.28, 14, NAVY, True)
    text(slide, body, x + 0.22, y + 0.57, width - 0.4, height - 0.7, 10, MUTED)


def bullet(slide, value, x, y, width, color=INK, size=11):
    shape(slide, MSO_SHAPE.OVAL, x, y + 0.08, 0.09, 0.09, CYAN)
    text(slide, value, x + 0.2, y, width - 0.2, 0.28, size, color)


def arrow(slide, x1, y1, x2, y2, color=BLUE):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = color
    connector.line.width = Pt(2)
    connector.line.end_arrowhead = True


# 1. Reference-style title slide
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.18, CYAN)
shape(slide, MSO_SHAPE.RECTANGLE, 0, 7.2, 13.333, 0.3, BLUE)
text(slide, "SMART INDIA HACKATHON 2026", 0.7, 0.82, 5.0, 0.28, 14, CYAN, True)
text(slide, "PRAMAN SETU", 0.68, 1.48, 7.3, 0.8, 42, WHITE, True, font="Aptos Display")
text(slide, "Digital Companion for Field Drug Testing", 0.72, 2.38, 7.7, 0.45, 22, LIGHT, True, font="Aptos Display")
text(slide, "A camera-led, offline-first workflow for calibrated presumptive results and tamper-evident field records.", 0.75, 3.08, 6.45, 0.8, 16, RGBColor(205, 221, 232))
tag(slide, "SIH26231", 0.75, 4.42, 1.35, ORANGE)
tag(slide, "MINISTRY OF HOME AFFAIRS", 2.25, 4.42, 2.55, GREEN)
tag(slide, "SOFTWARE", 5.0, 4.42, 1.1, CYAN)
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 8.55, 1.0, 3.45, 4.9, LIGHT, LIGHT)
text(slide, "REFERENCE\nCARD", 9.05, 1.43, 2.45, 0.55, 18, NAVY, True, align=PP_ALIGN.CENTER)
colors = [RGBColor(250, 250, 244), RGBColor(18, 18, 18), RED, GREEN, BLUE, RGBColor(142, 142, 142)]
for index, fill in enumerate(colors):
    shape(slide, MSO_SHAPE.RECTANGLE, 9.12 + (index % 3) * 0.76, 2.45 + (index // 3) * 0.76, 0.57, 0.57, fill, fill)
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 9.06, 4.35, 2.45, 0.72, NAVY, NAVY)
text(slide, "CARD VALID", 9.06, 4.55, 2.45, 0.22, 15, CYAN, True, align=PP_ALIGN.CENTER)
text(slide, "Problem Statement ID: SIH26231", 0.75, 6.42, 4.0, 0.24, 11, WHITE, True)
text(slide, "SIMULATED PROTOTYPE - PRESUMPTIVE RESULT ONLY", 7.9, 7.27, 4.3, 0.12, 7, GOLD, True, align=PP_ALIGN.RIGHT)

# 2. Proposed solution
slide = prs.slides.add_slide(blank)
header(slide, "PROPOSED SOLUTION", "From a subjective color reading to a documented field workflow.", 2)
text(slide, "Praman Setu works alongside existing colorimetric kits - no new hardware required.", 0.6, 1.52, 7.8, 0.35, 19, NAVY, True, font="Aptos Display")
card(slide, "CALIBRATED CAPTURE", "Camera overlay aligns the reference card and test kit. Six known patches correct lighting and white balance.", 0.62, 2.15, 3.0, 1.5, ORANGE)
card(slide, "EXPLAINABLE RESULT", "LAB color measurement, Delta E, confidence, quality score, timing gates, and explicit Inconclusive reasons.", 3.88, 2.15, 3.0, 1.5, GREEN)
card(slide, "TAMPER-EVIDENT RECORD", "Original image hash, timestamp, GPS, operator, kit metadata, hash chain, device signature, and QR/PDF export.", 7.14, 2.15, 3.0, 1.5, BLUE)
card(slide, "OFFLINE FIELD LOG", "SQLite-first records with PENDING, SYNCED, and FAILED states. Supervisor review remains read-only.", 10.4, 2.15, 2.3, 1.5, CYAN)
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.62, 4.35, 12.05, 1.35, NAVY, NAVY)
text(slide, "CAPTURE  ->  CALIBRATE  ->  CLASSIFY  ->  SIGN  ->  VERIFY", 1.0, 4.68, 11.3, 0.4, 22, WHITE, True, align=PP_ALIGN.CENTER, font="Aptos Display")
text(slide, "Output: Preliminary Positive / Preliminary Negative / Inconclusive - never a replacement for laboratory confirmation.", 1.0, 5.24, 11.3, 0.22, 11, RGBColor(205, 221, 232), align=PP_ALIGN.CENTER)
text(slide, "Innovation and uniqueness", 0.62, 6.08, 2.2, 0.25, 11, BLUE, True)
bullet(slide, "Reference-card calibration + quality gates", 2.55, 6.03, 2.8)
bullet(slide, "Cryptographic evidence identity", 5.48, 6.03, 2.45)
bullet(slide, "Offline-first + retest comparison", 8.0, 6.03, 2.75)
bullet(slide, "Simulation-safe demonstration", 10.82, 6.03, 1.75)

# 3. Technical approach
slide = prs.slides.add_slide(blank)
header(slide, "TECHNICAL APPROACH", "A mobile workflow around a versioned Reference DB and canonical evidence contract.", 3)
text(slide, "TECH STACK", 0.62, 1.45, 2.0, 0.25, 11, BLUE, True)
stack = [
    ("Frontend", "React Native", "Cross-platform mobile workflow"),
    ("Camera", "VisionCamera", "Capture, preview, QR/barcode path"),
    ("Image analysis", "OpenCV adapter", "Blur, glare, clipping, ROI checks"),
    ("Color science", "LAB + Delta E", "Gain correction and classification"),
    ("Storage", "SQLite", "Migrations + transactional offline queue"),
    ("Crypto", "Quick Crypto + Keystore/Keychain", "SHA-256, signature, public key"),
    ("Export", "HTML-to-PDF + QR SVG", "Evidence report and verification bundle"),
    ("Verification", "Plain HTML + JavaScript", "Hash, signature, chain checks"),
]
for index, (area, technology, purpose) in enumerate(stack):
    x = 0.62 + (index % 2) * 6.18
    y = 1.85 + (index // 2) * 0.98
    card(slide, area.upper(), f"{technology}\n{purpose}", x, y, 5.65, 0.78, [CYAN, ORANGE, GREEN, BLUE][index % 4])
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.62, 5.92, 12.0, 0.62, LIGHT, LIGHT)
text(slide, "Core contract", 0.9, 6.1, 1.2, 0.2, 11, NAVY, True)
text(slide, "Original image bytes -> SHA-256 -> canonical JSON -> record hash -> device signature -> JSON + image + PDF + QR", 2.25, 6.09, 9.95, 0.22, 11, MUTED, True)

# 4. Feasibility and viability
slide = prs.slides.add_slide(blank)
header(slide, "FEASIBILITY AND VIABILITY", "Prototype scope is intentionally narrow, testable, and safe to demonstrate.", 4)
feasibility = [
    ("Technical", "Uses existing smartphone camera and established mobile primitives: React Native, SQLite, native crypto, and OpenCV adapter.", CYAN),
    ("Operational", "Offline-first behavior preserves local evidence when field connectivity is unavailable; retry is explicit.", GREEN),
    ("Evidence", "Image hash, canonical record, previous hash, signature, and verification page make integrity visible.", BLUE),
    ("Safety", "Simulation-only materials, presumptive wording, hard quality gates, and laboratory-confirmation guidance.", ORANGE),
]
for index, (title, body, accent) in enumerate(feasibility):
    x = 0.7 + (index % 2) * 6.1
    y = 1.6 + (index // 2) * 1.88
    card(slide, title.upper(), body, x, y, 5.6, 1.52, accent)
text(slide, "Implementation boundary", 0.7, 5.55, 2.1, 0.25, 11, BLUE, True)
bullet(slide, "Bundled versioned Reference DB for the prototype", 2.7, 5.5, 3.25)
bullet(slide, "Mock sync adapter; no claim of production backend", 6.0, 5.5, 3.05)
bullet(slide, "Prototype signing; no claim of trusted timestamps or agency PKI", 9.2, 5.5, 3.2)
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 6.15, 11.9, 0.48, NAVY, NAVY)
text(slide, "Future production scope: backend, agency PKI, HSM, trusted timestamps, full encryption, CNN classifier, and multi-agency sync.", 0.95, 6.29, 11.4, 0.18, 10, WHITE, True, align=PP_ALIGN.CENTER)

# 5. Impact and benefits
slide = prs.slides.add_slide(blank)
header(slide, "IMPACT AND BENEFITS", "A stronger field record without changing the chemistry or the kit.", 5)
text(slide, "Praman Setu improves the evidence around a presumptive field test.", 0.65, 1.5, 7.0, 0.35, 21, NAVY, True, font="Aptos Display")
benefits = [
    ("CONSISTENCY", "Reference-card calibration and color-space measurement reduce subjective interpretation.", RED),
    ("ACCOUNTABILITY", "Operator, kit, batch, timing, location, image hash, signature, and chain travel together.", BLUE),
    ("CONTINUITY", "Offline records survive weak connectivity and can be synchronized or retried later.", GREEN),
    ("SUPERVISION", "Reviewers see the same read-only record and can request retest or laboratory confirmation.", ORANGE),
]
for index, (title, body, accent) in enumerate(benefits):
    x = 0.72 + (index % 2) * 6.1
    y = 2.2 + (index // 2) * 1.62
    card(slide, title, body, x, y, 5.65, 1.28, accent)
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.72, 5.65, 11.9, 0.68, NAVY, NAVY)
text(slide, "Presumptive result  !=  laboratory confirmation", 1.0, 5.82, 11.3, 0.25, 18, GOLD, True, align=PP_ALIGN.CENTER, font="Aptos Display")
text(slide, "The product is intentionally honest about uncertainty - that is part of its value.", 1.2, 6.48, 10.8, 0.22, 12, MUTED, align=PP_ALIGN.CENTER)

# 6. Research and references
slide = prs.slides.add_slide(blank)
header(slide, "RESEARCH AND REFERENCES", "Legal context, forensic guidance, and technical standards - clearly separated.", 6)
text(slide, "Government and legal context", 0.62, 1.42, 5.5, 0.25, 12, NAVY, True)
legal_sources = [
    ("SIH26231 problem statement", "https://sih.gov.in/sih2026PS"),
    ("Ministry of Home Affairs", "https://www.mha.gov.in/"),
    ("Narcotics Control Bureau", "https://ncb.gov.in/"),
    ("India Code legislation portal", "https://www.indiacode.gov.in/"),
    ("MHA: Three New Criminal Laws", "https://www.mha.gov.in/en/commoncontent/three-new-criminal-laws"),
    ("BNSS 2023 official text - Section 105", "https://www.mha.gov.in/sites/default/files/2024-02/BNSS_2023_English.pdf"),
    ("BSA 2023 official text - Section 63(4)", "https://www.indiacode.nic.in/bitstream/123456789/20098/1/a2023-47.pdf"),
]
for index, (label, url) in enumerate(legal_sources):
    y = 1.78 + index * 0.55
    text(slide, f"[{index + 1}]", 0.72, y, 0.34, 0.2, 10, BLUE, True)
    hyperlink_text(slide, label, url, 1.12, y, 4.85, 0.25, 9)
text(slide, "Forensic, laboratory, and security references", 6.72, 1.42, 5.8, 0.25, 12, NAVY, True)
technical = [
    ("NDPS Act 1985 - Section 52A reference", "https://dor.gov.in/sites/default/files/NDPSAct1985.pdf"),
    ("Supreme Court of India - case-law portal", "https://www.sci.gov.in/"),
    ("NIST Forensic Science", "https://www.nist.gov/forensic-science"),
    ("NIST OSAC Seized Drugs Process Map", "https://www.nist.gov/news-events/news/2022/11/osacs-seized-drugs-subcommittee-develops-process-map"),
    ("SWGDRUG Approved Recommendations", "https://www.swgdrug.org/approved.htm"),
    ("ISO/IEC 17025:2017", "https://www.iso.org/standard/66912.html"),
    ("NIST FIPS 180-4 / FIPS 186-5", "https://csrc.nist.gov/pubs/fips/180-4/upd1/final"),
]
for index, (label, url) in enumerate(technical):
    y = 1.78 + index * 0.55
    text(slide, f"[{index + 8}]", 6.82, y, 0.34, 0.2, 10, BLUE, True)
    hyperlink_text(slide, label, url, 7.22, y, 5.45, 0.25, 9)
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 5.82, 11.9, 0.7, LIGHT, LIGHT)
text(slide, "Important boundary", 0.98, 6.03, 1.4, 0.2, 10, NAVY, True)
text(slide, "Praman Setu supports documentation and traceability; it does not replace authorized statutory procedure, Magistrate certification, forensic confirmation, or legal review.", 2.55, 6.02, 9.65, 0.22, 10, MUTED, True)
text(slide, "All links are embedded as clickable hyperlinks in the editable deck.", 8.7, 6.72, 3.85, 0.2, 8, BLUE, True, align=PP_ALIGN.RIGHT)

prs.save(OUTPUT)
print(f"Created {OUTPUT} with {len(prs.slides)} slides")