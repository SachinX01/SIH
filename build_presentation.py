from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUTPUT = "Praman_Setu_SIH26231_Presentation.pptx"

BG = RGBColor(21, 27, 30)
PANEL = RGBColor(31, 40, 43)
PAPER = RGBColor(242, 237, 226)
WHITE = RGBColor(248, 249, 244)
MUTED = RGBColor(178, 190, 184)
MINT = RGBColor(124, 224, 188)
CORAL = RGBColor(255, 117, 99)
YELLOW = RGBColor(247, 201, 91)
BLUE = RGBColor(124, 185, 255)
INK = RGBColor(26, 34, 35)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def add_shape(slide, shape_type, x, y, width, height, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def add_text(slide, text, x, y, width, height, size=18, color=WHITE, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraphs = text.split("\n")
    for index, line in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return box


def add_rich_text(slide, segments, x, y, width, height, size=18, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for text, color, bold in segments:
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def base_slide(title, kicker=None, number=None):
    slide = prs.slides.add_slide(blank)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = BG
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.08, CORAL)
    if kicker:
        add_text(slide, kicker.upper(), 0.62, 0.34, 5.5, 0.25, 9, MINT, True)
    add_text(slide, title, 0.6, 0.62, 11.5, 0.65, 27, WHITE, True, font="Aptos Display")
    if number is not None:
        add_text(slide, f"0{number}", 12.05, 0.47, 0.65, 0.3, 11, MUTED, True, align=PP_ALIGN.RIGHT)
    return slide


def footer(slide, label="PRAMAN SETU / SIH26231"):
    add_text(slide, label, 0.62, 7.15, 4.5, 0.18, 8, MUTED, True)
    add_text(slide, "SIMULATED PROTOTYPE - PRESUMPTIVE RESULT ONLY", 8.0, 7.15, 4.7, 0.18, 8, CORAL, True, align=PP_ALIGN.RIGHT)


def pill(slide, text, x, y, width, color=MINT, text_color=BG):
    shape = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, 0.34, color, color, True)
    add_text(slide, text, x, y + 0.02, width, 0.25, 9, text_color, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return shape


def bullet_list(slide, items, x, y, width, size=15, color=WHITE, gap=0.42):
    for index, item in enumerate(items):
        yy = y + index * gap
        add_shape(slide, MSO_SHAPE.OVAL, x, yy + 0.1, 0.1, 0.1, CORAL)
        add_text(slide, item, x + 0.22, yy, width - 0.22, gap, size, color)


def card(slide, x, y, width, height, title, body, accent=MINT, number=None):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, height, PANEL, PANEL, True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.07, height, accent)
    if number:
        add_text(slide, number, x + 0.25, y + 0.2, 0.45, 0.3, 11, accent, True)
        tx = x + 0.78
        tw = width - 1.0
    else:
        tx = x + 0.28
        tw = width - 0.5
    add_text(slide, title, tx, y + 0.2, tw, 0.34, 15, WHITE, True)
    add_text(slide, body, tx, y + 0.64, tw, height - 0.8, 11, MUTED)


def flow_node(slide, label, sublabel, x, y, width, accent):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, 1.0, PANEL, PANEL, True)
    add_shape(slide, MSO_SHAPE.OVAL, x + 0.15, y + 0.18, 0.58, 0.58, accent)
    add_text(slide, label, x + 0.83, y + 0.16, width - 0.95, 0.3, 13, WHITE, True)
    add_text(slide, sublabel, x + 0.83, y + 0.51, width - 0.95, 0.26, 9, MUTED)


def arrow(slide, x1, y1, x2, y2, color=MINT):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = color
    connector.line.width = Pt(2)
    connector.line.end_arrowhead = True
    return connector


# 1. Title
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = BG
add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.1, CORAL)
add_shape(slide, MSO_SHAPE.ARC, 8.5, -0.8, 6.0, 6.0, BG, MINT)
add_shape(slide, MSO_SHAPE.RECTANGLE, 9.35, 1.12, 2.65, 4.2, PAPER, PAPER)
add_shape(slide, MSO_SHAPE.RECTANGLE, 9.68, 1.5, 1.98, 0.14, CORAL)
for idx, color in enumerate([RGBColor(250, 250, 245), RGBColor(25, 25, 25), CORAL, MINT, BLUE, RGBColor(135, 135, 135)]):
    add_shape(slide, MSO_SHAPE.RECTANGLE, 9.68 + (idx % 3) * 0.67, 2.0 + (idx // 3) * 0.67, 0.48, 0.48, color, color)
add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 9.65, 3.55, 2.05, 0.85, BG, BG, True)
add_text(slide, "CARD\nVALID", 9.65, 3.68, 2.05, 0.5, 17, MINT, True, align=PP_ALIGN.CENTER)
add_text(slide, "SMART INDIA HACKATHON 2026", 0.75, 0.8, 5.5, 0.35, 12, MINT, True)
add_text(slide, "Praman Setu", 0.72, 1.5, 7.7, 1.0, 42, WHITE, True, font="Aptos Display")
add_text(slide, "Digital companion for field drug testing", 0.76, 2.55, 7.6, 0.55, 23, PAPER, False, font="Aptos Display")
add_text(slide, "A camera-led, offline-first workflow for calibrated presumptive results and tamper-evident field records.", 0.78, 3.4, 6.9, 0.75, 16, MUTED)
pill(slide, "SIH26231", 0.78, 5.05, 1.35, CORAL, WHITE)
pill(slide, "MHA / SOFTWARE", 2.25, 5.05, 1.85, MINT, BG)
add_text(slide, "SIMULATED PROTOTYPE - NOT A DRUG TEST", 0.78, 6.65, 4.4, 0.25, 10, CORAL, True)
add_text(slide, "01", 12.05, 7.05, 0.6, 0.2, 9, MUTED, True, align=PP_ALIGN.RIGHT)

# 2. Problem
slide = base_slide("The field result is visible. The evidence is not.", "01 / The gap", 2)
add_text(slide, "Today, a color change is interpreted by eye -\nwith no consistent calibration and no verifiable record of the moment it happened.", 0.65, 1.55, 6.1, 1.0, 22, PAPER, True, font="Aptos Display")
card(slide, 7.25, 1.55, 2.55, 1.6, "SUBJECTIVE", "Different officers can read the same color differently.", CORAL, "01")
card(slide, 10.0, 1.55, 2.55, 1.6, "UNVERIFIED", "No camera evidence, calibrated reference, or chain of custody.", YELLOW, "02")
card(slide, 7.25, 3.45, 2.55, 1.6, "DISCONNECTED", "Kit batch, timing, place, and operator data are easy to lose.", BLUE, "03")
card(slide, 10.0, 3.45, 2.55, 1.6, "OFFLINE", "Field connectivity cannot be assumed at the point of testing.", MINT, "04")
add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 4.45, 5.9, 1.25, PAPER, PAPER, True)
add_text(slide, "The opportunity", 1.02, 4.7, 2.0, 0.25, 11, INK, True)
add_text(slide, "Turn one subjective color reading into a documented, repeatable field workflow.", 1.02, 5.03, 5.0, 0.42, 18, INK, True, font="Aptos Display")
footer(slide)

# 3. Solution
slide = base_slide("Praman Setu makes the test traceable.", "02 / The solution", 3)
add_text(slide, "One flow. Three layers of confidence.", 0.68, 1.4, 5.5, 0.4, 22, PAPER, True, font="Aptos Display")
card(slide, 0.72, 2.1, 3.75, 2.5, "CALIBRATE", "Use six reference patches and white-patch gain correction to normalize the captured image.", CORAL, "01")
card(slide, 4.8, 2.1, 3.75, 2.5, "CLASSIFY", "Measure the reaction region in LAB space, compare with Delta E, and surface confidence plus uncertainty.", MINT, "02")
card(slide, 8.88, 2.1, 3.75, 2.5, "PROVE", "Bind the image, metadata, timing, GPS, operator, hash, signature, and chain position into one record.", BLUE, "03")
add_text(slide, "Works with existing kits. No new hardware. Offline first.", 0.75, 5.35, 7.2, 0.45, 20, MINT, True, font="Aptos Display")
add_text(slide, "Output: a presumptive field-test result and supporting digital record - never a replacement for laboratory confirmation.", 0.75, 5.92, 10.8, 0.4, 13, MUTED)
footer(slide)

# 4. Workflow
slide = base_slide("From camera frame to verifiable record", "03 / The workflow", 4)
nodes = [
    ("CAPTURE", "Reference card + kit", CORAL),
    ("CALIBRATE", "Six patches + gain", YELLOW),
    ("MEASURE", "ROI + LAB + Delta E", MINT),
    ("RECORD", "Hash + chain + sign", BLUE),
    ("VERIFY", "JSON + image + QR", PAPER),
]
start_x = 0.6
for index, (label, sublabel, accent) in enumerate(nodes):
    x = start_x + index * 2.55
    flow_node(slide, label, sublabel, x, 2.15, 2.15, accent)
    add_text(slide, str(index + 1), x + 0.31, 2.33, 0.26, 0.2, 13, BG, True, align=PP_ALIGN.CENTER)
    if index < len(nodes) - 1:
        arrow(slide, x + 2.18, 2.65, x + 2.48, 2.65, MINT)
add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.15, 4.15, 11.0, 1.3, PANEL, PANEL, True)
add_text(slide, "Every transition is explicit", 1.55, 4.42, 2.8, 0.28, 12, MINT, True)
add_text(slide, "KIT VALIDATED  ->  TIMER RUNNING  ->  CAPTURE WINDOW  ->  CARD VALID  ->  RESULT READY  ->  SAVED  ->  EXPORTED", 1.55, 4.83, 10.0, 0.35, 16, WHITE, True, font="Aptos Display")
add_text(slide, "Hard gates stop the workflow; soft quality signals explain uncertainty.", 1.55, 5.32, 8.4, 0.25, 11, MUTED)
footer(slide)

# 5. Classification
slide = base_slide("The app does not guess. It measures, gates, and explains.", "04 / Classification", 5)
add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 1.45, 4.0, 4.9, PAPER, PAPER, True)
add_text(slide, "REFERENCE CARD", 1.05, 1.78, 2.8, 0.25, 11, INK, True)
for idx, color in enumerate([RGBColor(250, 250, 245), RGBColor(18, 18, 18), CORAL, MINT, BLUE, RGBColor(140, 140, 140)]):
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.05 + (idx % 3) * 0.92, 2.35 + (idx // 3) * 0.92, 0.72, 0.72, color, color)
add_text(slide, "Tap six known patches", 1.05, 4.34, 2.8, 0.3, 16, INK, True, font="Aptos Display")
add_text(slide, "White / black / red / green / blue / gray", 1.05, 4.8, 3.1, 0.3, 11, RGBColor(76, 84, 80))
add_text(slide, "CARD VALID", 1.05, 5.52, 2.2, 0.35, 17, RGBColor(24, 138, 103), True)
add_shape(slide, MSO.SHAPE.OVAL if False else MSO_SHAPE.OVAL, 3.55, 5.48, 0.34, 0.34, MINT)
add_text(slide, "1", 5.45, 1.75, 0.32, 0.25, 12, CORAL, True)
add_text(slide, "Calibrate", 5.85, 1.72, 1.8, 0.3, 18, WHITE, True)
add_text(slide, "White-patch gain correction\nReference-card validation\nCard version + residual Delta E", 5.45, 2.18, 6.6, 0.85, 14, MUTED)
add_text(slide, "2", 5.45, 3.35, 0.32, 0.25, 12, MINT, True)
add_text(slide, "Measure", 5.85, 3.32, 1.8, 0.3, 18, WHITE, True)
add_text(slide, "Sample ROI median RGB -> LAB\nDelta E against the Reference DB\nConfidence + uncertainty", 5.45, 3.78, 6.6, 0.85, 14, MUTED)
add_text(slide, "3", 5.45, 4.96, 0.32, 0.25, 12, BLUE, True)
add_text(slide, "Gate", 5.85, 4.93, 1.8, 0.3, 18, WHITE, True)
add_text(slide, "Blur / clipping / glare / ROI spread\nTiming window / card validity\nPreliminary result or Inconclusive", 5.45, 5.39, 6.6, 0.85, 14, MUTED)
footer(slide)

# 6. Integrity
slide = base_slide("Evidence is designed as a bundle, not a screenshot.", "05 / Integrity", 6)
add_text(slide, "The original image stays untouched. The analysis copy can change; the evidence identity cannot.", 0.7, 1.42, 8.7, 0.45, 19, PAPER, True, font="Aptos Display")
integrity = [
    ("ORIGINAL IMAGE", "SHA-256", CORAL),
    ("CANONICAL JSON", "record bytes", YELLOW),
    ("HASH CHAIN", "previous_hash", MINT),
    ("SIGNATURE", "device key + public key", BLUE),
    ("EXPORT", "JSON + image + PDF + QR", PAPER),
]
for index, (label, sublabel, accent) in enumerate(integrity):
    x = 0.72 + index * 2.48
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.35, 2.05, 1.55, PANEL, PANEL, True)
    add_shape(slide, MSO_SHAPE.OVAL, x + 0.78, 2.58, 0.5, 0.5, accent)
    add_text(slide, label, x + 0.18, 3.22, 1.7, 0.22, 10, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(slide, sublabel, x + 0.18, 3.5, 1.7, 0.18, 9, MUTED, align=PP_ALIGN.CENTER)
    if index < len(integrity) - 1:
        arrow(slide, x + 2.08, 3.12, x + 2.39, 3.12, MINT)
add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.3, 4.65, 10.7, 0.95, PAPER, PAPER, True)
add_text(slide, "VERIFY.HTML", 1.65, 4.95, 1.65, 0.22, 11, INK, True)
add_text(slide, "recomputes the image hash  |  checks the signature  |  walks the chain  |  reports each status separately", 3.3, 4.92, 8.25, 0.25, 13, INK, True)
add_text(slide, "Prototype limitation: this is tamper-evident evidence, not trusted timestamping or agency PKI.", 1.35, 5.95, 10.4, 0.3, 12, CORAL, True, align=PP_ALIGN.CENTER)
footer(slide)

# 7. Offline and roles
slide = base_slide("Built for the field: offline first, reviewable later.", "06 / Operations", 7)
card(slide, 0.72, 1.5, 3.7, 3.9, "OFFLINE QUEUE", "Every test is saved locally before synchronization.\n\nPENDING -> SYNCED\nPENDING -> FAILED\nFAILED -> RETRY\n\nLocal evidence is retained in every state.", MINT, "01")
card(slide, 4.82, 1.5, 3.7, 3.9, "SUPERVISOR REVIEW", "Read-only review of result, capture quality, kit, timing, metadata, and sync state.\n\nREVIEWED\nREQUIRES RETEST\nREQUIRES LAB CONFIRMATION", BLUE, "02")
card(slide, 8.92, 1.5, 3.7, 3.9, "RETEST, NOT MERGE", "Test 2 links to Test 1 through retest_of.\n\nEach record keeps its own image, hash, signature, and chain position.\n\nComparison is display-only.", CORAL, "03")
add_text(slide, "Field reality is a first-class requirement, not an edge case.", 0.76, 6.0, 8.2, 0.35, 20, PAPER, True, font="Aptos Display")
footer(slide)

# 8. Architecture
slide = base_slide("A thin mobile workflow around a strong evidence core.", "07 / Architecture", 8)
layers = [
    ("EXPERIENCE", "React Native screens, navigation, role routes", CORAL),
    ("WORKFLOW", "test state machine, timer, quality gates", YELLOW),
    ("DOMAIN", "calibration, LAB/Delta E, evidence builder", MINT),
    ("ADAPTERS", "VisionCamera, OpenCV, SQLite, GPS, keys, PDF", BLUE),
    ("VERIFY", "canonical JSON + image + signature checks", PAPER),
]
for index, (label, text, accent) in enumerate(layers):
    y = 1.45 + index * 0.87
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.86, y, 7.3, 0.63, PANEL, PANEL, True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.86, y, 0.16, 0.63, accent)
    add_text(slide, label, 1.25, y + 0.15, 1.4, 0.22, 11, accent, True)
    add_text(slide, text, 2.75, y + 0.14, 5.0, 0.25, 13, WHITE, True)
add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 9.0, 1.55, 3.2, 3.95, PAPER, PAPER, True)
add_text(slide, "PROTOTYPE\nBOUNDARY", 9.38, 1.95, 2.45, 0.7, 19, INK, True, font="Aptos Display", align=PP_ALIGN.CENTER)
bullet_list(slide, ["Bundled Reference DB", "Local SQLite", "Mock sync adapter", "Prototype signing", "No legal-evidence claim"], 9.35, 3.1, 2.35, 12, RGBColor(76, 84, 80), 0.43)
add_text(slide, "One canonical export contract connects the app to verify.html.", 0.9, 6.1, 8.0, 0.32, 17, MINT, True, font="Aptos Display")
footer(slide)

# 9. Demo
slide = base_slide("The demo proves the workflow without pretending to prove a drug.", "08 / Demonstration", 9)
add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 1.45, 4.0, 4.75, CORAL, CORAL, True)
add_text(slide, "SIMULATED\n- NOT A DRUG TEST", 1.05, 1.95, 3.3, 0.8, 23, WHITE, True, font="Aptos Display", align=PP_ALIGN.CENTER)
add_text(slide, "Use printed swatches, colored-water vials, simulated pouches, or printed rectangles. Never attach a simulated color to a named substance.", 1.1, 3.2, 3.2, 1.1, 15, WHITE, True, align=PP_ALIGN.CENTER)
add_text(slide, "The judge sees the same controls: QR validation, timer lock, card validation, quality gate, retest, verification, and review.", 1.1, 4.8, 3.2, 0.85, 13, WHITE, align=PP_ALIGN.CENTER)
add_text(slide, "JUDGE FLOW", 5.45, 1.55, 2.0, 0.25, 11, MINT, True)
demo_steps = [
    ("01", "Scan", "Valid kit QR continues; expired kit blocks."),
    ("02", "Capture", "Reference card + simulated sample."),
    ("03", "Classify", "Preliminary result with confidence."),
    ("04", "Prove", "Hash, signature, PDF, QR, verification."),
]
for index, (num, title, body) in enumerate(demo_steps):
    y = 2.0 + index * 0.9
    add_text(slide, num, 5.45, y, 0.42, 0.25, 12, CORAL, True)
    add_text(slide, title, 6.05, y, 1.4, 0.25, 16, WHITE, True)
    add_text(slide, body, 7.65, y, 4.6, 0.35, 12, MUTED)
add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 5.4, 5.75, 7.0, 0.55, PANEL, PANEL, True)
add_text(slide, "Every screen carries the presumptive-result disclaimer.", 5.65, 5.92, 6.5, 0.2, 12, PAPER, True, align=PP_ALIGN.CENTER)
footer(slide)

# 10. Impact
slide = base_slide("A stronger field record, without changing the field kit.", "09 / Impact", 10)
add_text(slide, "Praman Setu does not replace the chemistry. It improves the moment around it.", 0.72, 1.43, 8.5, 0.45, 21, PAPER, True, font="Aptos Display")
impact = [
    ("CONSISTENCY", "Calibrated image capture reduces subjective color interpretation.", CORAL),
    ("ACCOUNTABILITY", "Timestamp, GPS, operator, kit, timing, and image identity travel together.", MINT),
    ("CONTINUITY", "Offline records survive weak connectivity and can be reviewed later.", BLUE),
    ("CAUTION", "Inconclusive results and laboratory confirmation remain visible, never hidden.", YELLOW),
]
for index, (label, body, accent) in enumerate(impact):
    x = 0.75 + (index % 2) * 6.1
    y = 2.25 + (index // 2) * 1.85
    card(slide, x, y, 5.55, 1.42, label, body, accent)
add_text(slide, "Prototype deliverable: a working, explainable, judge-ready field workflow.", 0.78, 6.2, 9.8, 0.35, 18, MINT, True, font="Aptos Display")
footer(slide)

# 11. Close
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = BG
add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.1, CORAL)
add_text(slide, "Praman Setu", 0.8, 1.05, 8.4, 0.7, 39, WHITE, True, font="Aptos Display")
add_text(slide, "Make the presumptive result visible.\nMake the record verifiable.", 0.82, 2.05, 7.4, 1.1, 25, PAPER, True, font="Aptos Display")
add_text(slide, "SIH26231  /  Ministry of Home Affairs  /  Software", 0.84, 3.55, 6.5, 0.3, 13, MINT, True)
add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 8.8, 1.18, 3.45, 3.65, PAPER, PAPER, True)
add_text(slide, "CAPTURE", 9.25, 1.65, 2.5, 0.3, 16, INK, True, align=PP_ALIGN.CENTER)
add_text(slide, "+", 10.35, 2.2, 0.3, 0.4, 25, CORAL, True, align=PP_ALIGN.CENTER)
add_text(slide, "CALIBRATE", 9.25, 2.85, 2.5, 0.3, 16, INK, True, align=PP_ALIGN.CENTER)
add_text(slide, "+", 10.35, 3.4, 0.3, 0.4, 25, CORAL, True, align=PP_ALIGN.CENTER)
add_text(slide, "VERIFY", 9.25, 4.05, 2.5, 0.3, 16, INK, True, align=PP_ALIGN.CENTER)
add_text(slide, "SIMULATED PROTOTYPE - PRESUMPTIVE RESULT ONLY", 0.84, 6.7, 5.0, 0.25, 10, CORAL, True)
add_text(slide, "11", 12.05, 7.05, 0.6, 0.2, 9, MUTED, True, align=PP_ALIGN.RIGHT)

prs.save(OUTPUT)
print(f"Created {OUTPUT} with {len(prs.slides)} slides")